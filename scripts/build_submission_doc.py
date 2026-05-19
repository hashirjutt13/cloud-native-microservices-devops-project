from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Cloud_Computing_Final_Project_Submission.docx"
SHOT_DIR = ROOT / "docs" / "screenshots" / "team-lead"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(23, 32, 51)
MUTED = RGBColor(102, 112, 133)
LIGHT_FILL = "F2F4F7"
BORDER = "D9E1EC"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_para(doc, text="", style=None, before=0, after=6, align=None, bold=False, italic=False, color=INK, size=11):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=16 if level == 1 else 13, color=BLUE if level <= 2 else DARK_BLUE, bold=True)
    return p


def add_table(doc, rows, widths):
    table = doc.add_table(rows=1, cols=len(widths))
    table.autofit = False
    for idx, width in enumerate(widths):
        table.columns[idx].width = Inches(width)
    hdr = table.rows[0].cells
    for idx, value in enumerate(rows[0]):
        hdr[idx].text = str(value)
        set_cell_shading(hdr[idx], LIGHT_FILL)
        set_cell_border(hdr[idx])
        for p in hdr[idx].paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                set_run_font(r, bold=True, color=INK)
        hdr[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    for row in rows[1:]:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
            set_cell_border(cells[idx])
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cells[idx].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    set_run_font(r, color=INK)
    return table


def add_caption(doc, caption):
    p = add_para(doc, caption, before=2, after=8, size=9, color=MUTED, italic=True)
    return p


def add_image(doc, image_name, caption):
    path = SHOT_DIR / image_name
    add_heading(doc, caption, level=2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.35))
    add_caption(doc, image_name)


def add_placeholder(doc, label, note):
    add_heading(doc, label, level=2)
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.35)
    cell = table.cell(0, 0)
    set_cell_border(cell, color="7F8EA3", size="10")
    set_cell_shading(cell, "FBFDFF")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    row = table.rows[0]
    row.height = Inches(2.7)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(26)
    r = p.add_run("PASTE SCREENSHOT HERE\\n")
    set_run_font(r, size=14, color=MUTED, bold=True)
    r2 = p.add_run(note)
    set_run_font(r2, size=10, color=MUTED)
    add_para(doc, "", after=8)


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color in [
        ("Title", 23, RGBColor(0, 0, 0)),
        ("Subtitle", 14, MUTED),
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DARK_BLUE),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = name.startswith("Heading")
        style.paragraph_format.space_before = Pt(16 if name == "Heading 1" else 12)
        style.paragraph_format.space_after = Pt(8 if name == "Heading 1" else 6)
        style.paragraph_format.line_spacing = 1.10


def add_footer(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Cloud Computing Final Project | Hashir Sarwar")
    set_run_font(r, size=9, color=MUTED)


def main():
    doc = Document()
    configure_styles(doc)
    add_footer(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Cloud Computing Final Project Submission")
    set_run_font(r, size=23, bold=True, color=RGBColor(0, 0, 0))

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Cloud-Native Microservices DevOps Project")
    set_run_font(r, size=14, color=MUTED)

    add_para(doc, "Course: Lab-DevOps for Cloud Computing", align=WD_ALIGN_PARAGRAPH.CENTER, color=MUTED)
    add_para(doc, "Repository: https://github.com/hashirjutt13/cloud-native-microservices-devops-project", align=WD_ALIGN_PARAGRAPH.CENTER, color=MUTED)

    add_heading(doc, "Team Members", level=1)
    add_table(
        doc,
        [
            ["Name", "Roll Number", "Role", "Assigned Area"],
            ["Hashir Sarwar", "fa23-bcs-065", "Team Lead", "Dashboard, repo setup, CI/CD, final submission"],
            ["Soban Rabbani", "fa23-bcs-082", "Member", "User/Product contribution and PR proof"],
            ["Abdul Hadi", "fa23-bcs-012", "Member", "Order/Notification contribution and PR proof"],
        ],
        [1.55, 1.25, 1.2, 2.35],
    )

    add_heading(doc, "Project Summary", level=1)
    add_para(
        doc,
        "This project implements a DevOps-focused cloud-native e-commerce platform with a frontend dashboard, four REST microservices, Docker containerization, GitHub Actions, Jenkins pipeline artifacts, Kubernetes manifests, Render deployment hooks, and documentation for final submission.",
    )
    add_table(
        doc,
        [
            ["Item", "Status"],
            ["GitHub repository", "Created and public"],
            ["Git Flow branches", "main, develop, release"],
            ["GitHub Environments", "development, staging, production"],
            ["Docker Hub", "Images published under hashirsarwar"],
            ["GitHub Actions", "CI, Docker Publish, and Render Deploy successful"],
            ["Team member PRs", "Pending: Soban and Abdul Hadi"],
        ],
        [2.2, 4.15],
    )

    add_heading(doc, "Submission URLs", level=1)
    add_table(
        doc,
        [
            ["Field", "URL / Value"],
            ["GitHub Repository", "https://github.com/hashirjutt13/cloud-native-microservices-devops-project"],
            ["Docker Hub", "https://hub.docker.com/u/hashirsarwar"],
            ["Development Environment", "Paste final Render development URL here"],
            ["QA/Staging Environment", "Paste final Render staging URL here"],
            ["Production Environment", "Paste final Render production URL here"],
        ],
        [2.2, 4.15],
    )

    doc.add_page_break()
    add_heading(doc, "Team Lead Screenshots", level=1)
    add_para(doc, "The following screenshots were captured from the team-lead side and are ready for submission evidence.")

    screenshots = [
        ("teamlead-01-repo-home.png", "GitHub Repository Home"),
        ("teamlead-02-branches.png", "GitHub Branches"),
        ("teamlead-03-actions.png", "GitHub Actions Runs"),
        ("teamlead-04-issues.png", "GitHub Issues And Member Tasks"),
        ("teamlead-05-dockerhub-images.png", "Docker Hub Images"),
        ("teamlead-13-setup-proof.png", "Repository Setup Proof"),
        ("teamlead-08-local-dashboard.png", "Local Dashboard"),
        ("teamlead-09-local-users.png", "Local Users Page"),
        ("teamlead-10-local-products.png", "Local Products Page"),
        ("teamlead-11-local-orders.png", "Local Orders Page"),
        ("teamlead-12-local-notifications.png", "Local Notifications Page"),
    ]
    for image, caption in screenshots:
        add_image(doc, image, caption)

    doc.add_page_break()
    add_heading(doc, "Pending Team Member Screenshots", level=1)
    add_para(
        doc,
        "Paste the following screenshots after Soban Rabbani and Abdul Hadi complete their feature branches and pull requests. These placeholders are intentionally blank.",
    )

    soban = [
        ("Soban Screenshot 1: Repo Access", "Show the repository page after accepting the invitation."),
        ("Soban Screenshot 2: Local Checks", "Show terminal output after npm run check passes."),
        ("Soban Screenshot 3: Feature Branch", "Show branch feature/soban-user-product on GitHub."),
        ("Soban Screenshot 4: Commit Proof", "Show Soban's commit author and message."),
        ("Soban Screenshot 5: Pull Request", "Show PR title and base branch develop."),
        ("Soban Screenshot 6: PR CI Passed", "Show green checks on the PR."),
    ]
    abdul = [
        ("Abdul Hadi Screenshot 1: Repo Access", "Show the repository page after accepting the invitation."),
        ("Abdul Hadi Screenshot 2: Local Checks", "Show terminal output after npm run check passes."),
        ("Abdul Hadi Screenshot 3: Feature Branch", "Show branch feature/abdul-order-notification on GitHub."),
        ("Abdul Hadi Screenshot 4: Commit Proof", "Show Abdul Hadi's commit author and message."),
        ("Abdul Hadi Screenshot 5: Pull Request", "Show PR title and base branch develop."),
        ("Abdul Hadi Screenshot 6: PR CI Passed", "Show green checks on the PR."),
    ]

    for label, note in soban:
        add_placeholder(doc, label, note)
    for label, note in abdul:
        add_placeholder(doc, label, note)

    add_heading(doc, "Reflection Drafts", level=1)
    add_para(doc, "Hashir Sarwar: I learned how a DevOps project is structured end to end, from Git Flow and branch protection to Docker images, CI/CD pipelines, and Kubernetes deployment.")
    add_para(doc, "Soban Rabbani: I worked with service/page ownership and saw how individual contributions fit into a shared repository through feature branches and pull requests.")
    add_para(doc, "Abdul Hadi: I focused on deployment and pipeline artifacts, including Jenkins, Kubernetes manifests, and environment promotion.")

    add_heading(doc, "Final Checklist", level=1)
    checklist = [
        "Soban accepts GitHub invitation and opens PR into develop.",
        "Abdul Hadi accepts GitHub invitation and opens PR into develop.",
        "Team lead merges PRs and promotes develop to release, then release to main.",
        "Paste remaining member screenshots into this document.",
        "Paste final Render environment URLs.",
        "Export or submit this Word document according to instructor instructions.",
    ]
    for item in checklist:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
